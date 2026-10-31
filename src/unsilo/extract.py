"""Turn file bytes (PDF, Word, Excel, CSV, text) into plain text for classification and search."""
from __future__ import annotations

import csv
import io

MAX_CHARS = 60_000  # enough for any sandbox document; keeps prompts bounded

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV = "text/csv"


def extract_text(data: bytes, mime_type: str, filename: str = "") -> str:
    name = filename.lower()
    if mime_type == PDF or name.endswith(".pdf"):
        return _pdf(data)
    if mime_type == DOCX or name.endswith(".docx"):
        return _docx(data)
    if mime_type == XLSX or name.endswith((".xlsx", ".xlsm")):
        return _xlsx(data)
    if mime_type == CSV or name.endswith(".csv"):
        return _csv(data)
    if mime_type.startswith("text/") or name.endswith((".txt", ".md", ".json")):
        return data.decode("utf-8", errors="replace")[:MAX_CHARS]
    raise ValueError(f"unsupported file type: {mime_type or filename}")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)[:MAX_CHARS]


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)[:MAX_CHARS]


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                lines.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)[:MAX_CHARS]


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    rows = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(r) for r in rows)[:MAX_CHARS]
