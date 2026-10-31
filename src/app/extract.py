"""File bytes -> plain text (PDF, Word, Excel, CSV, text)."""
import csv
import io

LIMIT = 60_000
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def to_text(data: bytes, mime: str, name: str = "") -> str:
    n = name.lower()
    if mime == PDF or n.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)[:LIMIT]
    if mime == DOCX or n.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        lines += [" | ".join(c.text.strip() for c in row.cells) for t in d.tables for row in t.rows]
        return "\n".join(lines)[:LIMIT]
    if mime == XLSX or n.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# Sheet: {ws.title}")
            lines += [" | ".join("" if v is None else str(v) for v in row)
                      for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)]
        return "\n".join(lines)[:LIMIT]
    if mime == "text/csv" or n.endswith(".csv"):
        return "\n".join(" | ".join(r) for r in csv.reader(io.StringIO(data.decode("utf-8", "replace"))))[:LIMIT]
    if mime.startswith("text/") or n.endswith((".txt", ".md", ".json")):
        return data.decode("utf-8", "replace")[:LIMIT]
    raise ValueError(f"unsupported type {mime or name}")
